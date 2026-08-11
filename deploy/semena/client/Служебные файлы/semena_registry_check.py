from __future__ import annotations

import argparse
import concurrent.futures
import html
import pathlib
import re
import sys
import time
import urllib.parse
import urllib.request
from typing import Iterable

import openpyxl
import xlrd


REGISTRY_URL = (
    "https://gossortrf.ru/registry/"
    "gosudarstvennyy-reestr-selektsionnykh-dostizheniy-dopushchennykh-k-ispolzovaniyu-tom-1-sorta-rasteni/"
)
USER_AGENT = "Semena-Agent registry checker/1.0"
NOT_FOUND = "нету"


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize(value: object) -> str:
    text = clean(value).lower().replace("ё", "е")
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"\[[^\]]*\]", " ", text)
    text = re.sub(r"(?i)\bf\s*1\b", "f1", text)
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def variants(variety: str) -> list[str]:
    result: list[str] = []

    def add(value: str) -> None:
        value = re.sub(r"\s+", " ", clean(value)).strip(" ,;")
        if value and value not in result:
            result.append(value)

    add(variety)
    add(re.sub(r"\([^)]*\)", " ", variety))
    add(re.sub(r"\[[^\]]*\]", " ", variety))
    no_parentheses = re.sub(r"\([^)]*\)", " ", variety)
    add(re.split(r"[,;]", no_parentheses)[0])
    add(re.sub(r"(?i)\bf\s*1\b", "F1", no_parentheses))
    return result[:4]


def request_registry(culture: str, variety: str) -> str:
    query = urllib.parse.urlencode(
        {
            "arrFilter_pf[CULTURE_NAME]": culture,
            "arrFilter_pf[SORT_NAME]": variety,
            "set_filter": "Y",
        }
    )
    request = urllib.request.Request(
        REGISTRY_URL + "?" + query,
        headers={"User-Agent": USER_AGENT},
    )
    with urllib.request.urlopen(request, timeout=25) as response:
        charset = response.headers.get_content_charset() or "utf-8"
        return response.read().decode(charset, errors="replace")


def strip_tags(value: str) -> str:
    return clean(html.unescape(re.sub(r"<[^>]+>", " ", value)))


def first_class_text(block: str, class_name: str) -> str:
    pattern = re.compile(
        r'<(?P<tag>span|div)[^>]*class="[^"]*\b'
        + re.escape(class_name)
        + r'\b[^"]*"[^>]*>(?P<body>.*?)</(?P=tag)>',
        re.IGNORECASE | re.DOTALL,
    )
    match = pattern.search(block)
    return strip_tags(match.group("body")) if match else ""


def parse_year(page: str, culture: str, wanted: str) -> str | None:
    if "Найдено: 0" in strip_tags(page):
        return None

    culture_norm = normalize(culture)
    wanted_norm = normalize(wanted)
    candidates: list[tuple[str, str, str]] = []
    for match in re.finditer(
        r'<a[^>]*class="[^"]*\bregistry__results-table-link\b[^"]*"[^>]*>(.*?)</a>',
        page,
        re.IGNORECASE | re.DOTALL,
    ):
        block = match.group(1)
        row_culture = first_class_text(block, "registry__results-table-cell__culture").split("(")[0].strip()
        row_name = first_class_text(block, "registry__results-table-cell__name")
        row_year = first_class_text(block, "registry__results-table-cell__year")
        if re.fullmatch(r"\d{4}", row_year):
            candidates.append((normalize(row_culture), normalize(row_name), row_year))

    for row_culture, row_name, row_year in candidates:
        if row_culture == culture_norm and row_name == wanted_norm:
            return row_year
    for _row_culture, row_name, row_year in candidates:
        if row_name == wanted_norm:
            return row_year

    match = re.search(r"Год включения:\s*(\d{4})", strip_tags(page))
    return match.group(1) if match else None


def lookup(pair: tuple[str, str]) -> tuple[tuple[str, str], str]:
    culture, variety = pair
    for candidate in variants(variety):
        try:
            year = parse_year(request_registry(culture, candidate), culture, candidate)
            if year:
                return pair, year
        except Exception:
            pass
    return pair, NOT_FOUND


def find_source(workspace: pathlib.Path, requested: str) -> pathlib.Path:
    direct = workspace / requested
    if direct.is_file():
        return direct

    requested_norm = normalize(pathlib.Path(requested).stem)
    candidates = [
        path
        for path in workspace.glob("*")
        if path.is_file() and path.suffix.lower() in {".xls", ".xlsx", ".xlsm"}
    ]
    scored = []
    for path in candidates:
        name_norm = normalize(path.stem)
        score = sum(1 for part in requested_norm.split() if part in name_norm)
        if requested_norm and requested_norm in name_norm:
            score += 10
        scored.append((score, path))
    scored.sort(key=lambda item: (item[0], -len(item[1].name)), reverse=True)
    if scored and scored[0][0] > 0:
        return scored[0][1]
    raise FileNotFoundError(f"Excel-файл не найден: {requested}")


def read_xls(path: pathlib.Path) -> list[list[object]]:
    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    rows: list[list[object]] = []
    for row_index in range(sheet.nrows):
        row: list[object] = []
        for col_index in range(sheet.ncols):
            value = sheet.cell_value(row_index, col_index)
            row.append("" if value == "" else value)
        rows.append(row)
    return rows


def read_xlsx(path: pathlib.Path) -> list[list[object]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def extract_rows(rows: list[list[object]]) -> list[tuple[int, str, str]]:
    result: list[tuple[int, str, str]] = []
    for index, row in enumerate(rows):
        if len(row) <= 10:
            continue
        number = clean(row[1])
        culture = clean(row[4])
        variety = clean(row[10])
        if not number.isdigit() or not culture or not variety:
            continue
        if re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}", culture):
            continue
        result.append((index, culture, variety))
    return result


def unique_pairs(items: Iterable[tuple[int, str, str]]) -> list[tuple[str, str]]:
    seen: set[tuple[str, str]] = set()
    result: list[tuple[str, str]] = []
    for _index, culture, variety in items:
        key = (normalize(culture), normalize(variety))
        if key not in seen:
            seen.add(key)
            result.append((culture, variety))
    return result


def write_result(source_rows: list[list[object]], item_rows: list[tuple[int, str, str]], years: dict[tuple[str, str], str], output: pathlib.Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Проверка"
    for row_index, row in enumerate(source_rows, 1):
        for col_index, value in enumerate(row, 1):
            if value != "":
                sheet.cell(row_index, col_index).value = value

    result_col = max((len(row) for row in source_rows), default=0) + 1
    header_row = 78 if len(source_rows) >= 78 else 1
    sheet.cell(header_row, result_col).value = "Год включения в Госреестр (gossortrf.ru)"
    for row_index, culture, variety in item_rows:
        sheet.cell(row_index + 1, result_col).value = years.get((normalize(culture), normalize(variety)), NOT_FOUND)
    workbook.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Проверка Excel-файла по реестру Госсорткомиссии")
    parser.add_argument("excel", help="Имя Excel-файла в рабочей папке")
    parser.add_argument("--workers", type=int, default=18)
    args = parser.parse_args()

    workspace = pathlib.Path.cwd()
    source = find_source(workspace, args.excel)
    output = source.with_name(source.stem + " - проверка через сайт госреестра.xlsx")
    print(f"Файл: {source}")

    source_rows = read_xls(source) if source.suffix.lower() == ".xls" else read_xlsx(source)
    item_rows = extract_rows(source_rows)
    pairs = unique_pairs(item_rows)
    print(f"Строк к проверке: {len(item_rows)}")
    print(f"Уникальных пар культура/сорт: {len(pairs)}")

    started = time.time()
    years: dict[tuple[str, str], str] = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {executor.submit(lookup, pair): pair for pair in pairs}
        for done, future in enumerate(concurrent.futures.as_completed(futures), 1):
            pair, year = future.result()
            years[(normalize(pair[0]), normalize(pair[1]))] = year
            if done <= 5 or done % 100 == 0 or done == len(pairs):
                found = sum(1 for value in years.values() if value != NOT_FOUND)
                print(f"Проверено {done}/{len(pairs)}, найдено {found}, секунд {time.time() - started:.1f}")

    write_result(source_rows, item_rows, years, output)
    found_rows = sum(
        1
        for _index, culture, variety in item_rows
        if years.get((normalize(culture), normalize(variety))) != NOT_FOUND
    )
    print(f"Готово: {output}")
    print(f"Обработано строк: {len(item_rows)}")
    print(f"Найдено строк: {found_rows}")
    print(f"Не найдено строк: {len(item_rows) - found_rows}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
